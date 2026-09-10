import os
from pptx import Presentation

pptx_path = r'C:\Proyectos\Interactivo\Interactivo\ES\DDEE BRIG UU PPUU.pptx'
if not os.path.exists(pptx_path):
    print("Not found:", pptx_path)
else:
    prs = Presentation(pptx_path)
    print(f"Total slides in DDEE BRIG UU PPUU.pptx: {len(prs.slides)}")
    
    # Check slide 18 (index 17) which corresponds to BIM 5 Zarumilla
    slide = prs.slides[17] # 18th slide (0-indexed 17)
    print(f"\n--- SLIDE 18 ---")
    for shape in slide.shapes:
        if shape.has_text_frame:
            print(f"Shape text: {repr(shape.text_frame.text)}")
