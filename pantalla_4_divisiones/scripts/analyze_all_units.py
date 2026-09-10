import os
import json
import re
from pptx import Presentation

pptx_path = r'C:\Proyectos\Interactivo\Interactivo\ES\DDEE BRIG UU PPUU.pptx'
prs = Presentation(pptx_path)

with open(r'c:\Proyectos\Interactivo\mvp_ejercito\pantalla_4_divisiones\src\data\unit_database.json', 'r', encoding='utf-8') as f:
    unit_db = json.load(f)

with open(r'c:\Proyectos\Interactivo\mvp_ejercito\pantalla_4_divisiones\src\data\divisiones.json', 'r', encoding='utf-8') as f:
    divs_data = json.load(f)

def clean_str(s):
    if not s: return ""
    s = s.replace('“', '"').replace('”', '"').replace('‘', "'").replace('’', "'")
    s = re.sub(r'\s+', ' ', s).strip()
    return s

def norm(s):
    s = clean_str(s).lower()
    return re.sub(r'[^a-z0-9]', '', s)

# 1. Map PPTX slides
pptx_slides = []
for idx, slide in enumerate(prs.slides):
    slide_num = idx + 1
    shapes_text = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            txt = clean_str(shape.text_frame.text)
            if txt:
                shapes_text.append(txt)
    
    if len(shapes_text) >= 2:
        body = max(shapes_text, key=len)
        if len(body) >= 80:
            titles = [t for t in shapes_text if t != body]
            title = " ".join(titles)
            pptx_slides.append({
                "slide_num": slide_num,
                "title": clean_str(title),
                "body": clean_str(body)
            })

# 2. Extract every unit clickable in divisiones.json
organigram_units = []
for div in divs_data['divisiones']:
    div_id = div.get('id')
    div_nombre = div.get('nombre')
    
    # Division level direct units
    for u in div.get('unidadesDivisionarias', []):
        organigram_units.append({
            "div_id": div_id,
            "div_nombre": div_nombre,
            "brigada": "Unidad Divisionaria Directa",
            "nombre": u
        })
    
    # Brigade level
    for b in div.get('brigadas', []):
        b_nombre = b.get('nombre')
        for u in b.get('unidades', []):
            u_nom = u if isinstance(u, str) else u.get('nombre', '')
            organigram_units.append({
                "div_id": div_id,
                "div_nombre": div_nombre,
                "brigada": b_nombre,
                "nombre": u_nom
            })

print(f"Total unidades en el organigrama de la app (divisiones.json): {len(organigram_units)}")
print(f"Total diapositivas con reseña en el PPTX: {len(pptx_slides)}")

# Now for each unit in the organigram, find what the app displays in the modal vs what the PPTX slide says
results = []
for unit in organigram_units:
    u_name = unit["nombre"]
    norm_u = norm(u_name)
    
    # How app resolves unit:
    db_entry = unit_db.get(norm_u)
    if not db_entry:
        # try find in unit_db keys
        for k, v in unit_db.items():
            if norm(k) == norm_u:
                db_entry = v
                break
    
    app_display_name = clean_str(db_entry.get("nombre") if db_entry else u_name)
    app_resena = clean_str(db_entry.get("resena", "") if db_entry else "")
    
    # Match with PPTX slide
    matched_slide = None
    if app_resena and len(app_resena) > 40:
        norm_res = norm(app_resena[:60])
        for s in pptx_slides:
            norm_s_body = norm(s["body"][:60])
            if norm_s_body.startswith(norm_res[:35]) or norm_res.startswith(norm_s_body[:35]):
                matched_slide = s
                break
    
    if not matched_slide:
        # Try match by title/number in slide
        num_match = re.search(r'n[°º\s]*(\d+)', u_name, re.I)
        if num_match:
            num = num_match.group(1)
            for s in pptx_slides:
                s_num_match = re.search(r'n[°º\s]*' + num + r'\b', s["title"], re.I)
                if s_num_match:
                    # check type overlap
                    if any(w in s["title"].upper() for w in ["INFANTERÍA", "CABALLERÍA", "ARTILLERÍA", "INGENIERÍA", "COMUNICACIONES", "SERVICIOS", "COMANDO", "POLICÍA"]):
                        matched_slide = s
                        break
                        
    if matched_slide:
        pptx_title = clean_str(matched_slide["title"])
        is_diff = norm(pptx_title) != norm(app_display_name)
        results.append({
            "div": unit["div_nombre"],
            "brigada": unit["brigada"],
            "organigrama_name": u_name,
            "app_modal_title": app_display_name,
            "pptx_slide_num": matched_slide["slide_num"],
            "pptx_slide_title": pptx_title,
            "different": is_diff
        })
    else:
        results.append({
            "div": unit["div_nombre"],
            "brigada": unit["brigada"],
            "organigrama_name": u_name,
            "app_modal_title": app_display_name,
            "pptx_slide_num": None,
            "pptx_slide_title": "NO ENCONTRADA EN PPTX",
            "different": True
        })

diff_results = [r for r in results if r["different"]]
same_results = [r for r in results if not r["different"]]

print(f"\nResultados del cruce:")
print(f"Total unidades analizadas: {len(results)}")
print(f"Unidades con DIFERENCIA entre Diapositiva y App: {len(diff_results)}")
print(f"Unidades con Título IDÉNTICO: {len(same_results)}")

with open(r'c:\Proyectos\Interactivo\mvp_ejercito\pantalla_4_divisiones\scripts\diferencias_detalle.json', 'w', encoding='utf-8') as f:
    json.dump(results, f, indent=2, ensure_ascii=False)
