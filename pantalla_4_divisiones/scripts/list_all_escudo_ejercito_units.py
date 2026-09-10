import json
import os
import re
from pptx import Presentation

base_dir = r'c:\Proyectos\Interactivo\mvp_ejercito\pantalla_4_divisiones'
db_path = os.path.join(base_dir, 'src', 'data', 'unit_database.json')
with open(db_path, 'r', encoding='utf-8') as f:
    unit_db = json.load(f)

# Find all entries with escudo_ejercito
escudo_ejercito_entries = {}
for k, v in unit_db.items():
    if 'escudo_ejercito' in v.get('escudo', ''):
        name = v.get('nombre', '').strip()
        if name:
            if name not in escudo_ejercito_entries:
                escudo_ejercito_entries[name] = v

print(f"Total distinct unit names with escudo_ejercito: {len(escudo_ejercito_entries)}")

pptx_path = r'C:\Proyectos\Interactivo\Interactivo\ES\DDEE BRIG UU PPUU.pptx'
prs = Presentation(pptx_path)

# Build slide index
slides_data = []
for idx, slide in enumerate(prs.slides):
    stext = ' '.join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    has_pic = any(s.shape_type == 13 for s in slide.shapes)
    pic_names = [s.name for s in slide.shapes if s.shape_type == 13]
    slides_data.append({
        'num': idx + 1,
        'text': stext,
        'text_lower': stext.lower(),
        'has_pic': has_pic,
        'pic_names': pic_names
    })

categorized = {
    'with_pic_in_pptx': [],
    'no_pic_in_pptx': [],
    'no_slide_found': []
}

for name, unit_obj in sorted(escudo_ejercito_entries.items()):
    clean = re.sub(r'[«»“”"\'\.]', '', name).strip()
    words = [w for w in clean.split() if len(w) > 3 and w.upper() not in [
        'BATALLÓN', 'COMPAÑÍA', 'REGIMIENTO', 'GRUPO', 'ESCUADRÓN', 'EJÉRCITO', 'UNIDAD',
        'SERVICIOS', 'COMUNICACIONES', 'INGENIERÍA', 'INFANTERÍA', 'CABALLERÍA', 'ARTILLERÍA', 'BLINDADO', 'MOTORIZADO'
    ]]
    nums = re.findall(r'\b\d{1,3}\b', clean)

    matched = []
    for s in slides_data:
        # Check if words match
        match_count = sum(1 for w in words if w.lower() in s['text_lower'])
        num_match = (not nums) or any(re.search(rf'\b{n}\b', s['text_lower']) for n in nums)
        if words and match_count >= min(2, len(words)) and num_match:
            matched.append(s)
        elif len(words) == 1 and match_count == 1 and num_match:
            matched.append(s)

    if matched:
        # Check if any matched slide has pictures
        has_any_pic = any(s['has_pic'] for s in matched)
        slide_nums = [str(s['num']) for s in matched]
        if has_any_pic:
            categorized['with_pic_in_pptx'].append({
                'name': name,
                'slides': slide_nums,
                'has_pic': True
            })
        else:
            categorized['no_pic_in_pptx'].append({
                'name': name,
                'slides': slide_nums,
                'has_pic': False
            })
    else:
        categorized['no_slide_found'].append({
            'name': name
        })

print(f"\n1. In PPTX with picture ({len(categorized['with_pic_in_pptx'])}):")
for item in categorized['with_pic_in_pptx']:
    print(f"  - {item['name']} (Slides: {', '.join(item['slides'])})")

print(f"\n2. In PPTX WITHOUT picture ({len(categorized['no_pic_in_pptx'])}):")
for item in categorized['no_pic_in_pptx']:
    print(f"  - {item['name']} (Slides: {', '.join(item['slides'])})")

print(f"\n3. No direct slide match found / Generic name ({len(categorized['no_slide_found'])}):")
for item in categorized['no_slide_found'][:40]:
    print(f"  - {item['name']}")
if len(categorized['no_slide_found']) > 40:
    print(f"  ... and {len(categorized['no_slide_found']) - 40} more")
