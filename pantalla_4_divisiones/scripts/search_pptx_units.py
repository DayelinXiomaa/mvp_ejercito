from pptx import Presentation

prs = Presentation(r'C:\Proyectos\Interactivo\Interactivo\ES\DDEE BRIG UU PPUU.pptx')

terms = ['espinar', 'colina', 'larico', 'chocano']

for idx, slide in enumerate(prs.slides):
    slide_text = []
    for s in slide.shapes:
        if s.has_text_frame:
            slide_text.append(s.text_frame.text)
    full_txt = ' '.join(slide_text).lower()
    for term in terms:
        if term in full_txt:
            has_img = any(s.shape_type == 13 for s in slide.shapes)
            print(f'Match "{term}" in Slide {idx+1}: has_picture={has_img}')
            for s in slide.shapes:
                if s.has_text_frame and term in s.text_frame.text.lower():
                    print('  Text:', repr(s.text_frame.text.replace('\n', ' ')[:90]))
                if s.shape_type == 13:
                    print('  Picture:', s.name, s.image.ext, len(s.image.blob), 'bytes')
